"""Bounded, disposable parser process for the file-extraction experiment.

Run this file directly, never in the web process: ``python <script> PATH EXT``.
PATH is a regular file staged by the trusted endpoint, not a URL or user path.
Stdout contains one bounded JSON object. File contents are always untrusted data.
The limits and Python audit guard are defense in depth, not OS/container isolation.
PDF/image conversion needs separately provisioned Docling and offline artifacts;
this module never downloads models or silently substitutes text-only PDF parsing.
"""

from __future__ import annotations

import json
import math
import os
import stat
import sys
from pathlib import Path
from typing import Any

MAX_INPUT_BYTES = 20 * 1024 * 1024
MAX_PAGES = 100
MAX_IMAGE_PIXELS = 20_000_000
OCR_RASTER_SCALE = 4.5
MAX_SHEETS = 32
MAX_CELLS = 100_000
MAX_COLUMNS = 1024
MAX_RECORDS = 20_000
MAX_RECORD_TEXT = 16_000
MAX_SECTIONS = 512
MAX_INDEX_ENTRIES = 512
MAX_INDEX_TEXT = 255
MAX_RESULT_BYTES = 8 * 1024 * 1024
MAX_ARCHIVE_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 2048
MAX_PARSER_SECONDS = 120
MAX_MEMORY_BYTES = 4 * 1024 * 1024 * 1024
IMAGE_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"})
SUPPORTED_EXTENSIONS = frozenset({"txt", "xlsx", "pdf"}) | IMAGE_EXTENSIONS
DOCLING_ARTIFACTS_ENV = "BLOOMERP_FILE_EXTRACTION_DOCLING_ARTIFACTS"


class ParseError(Exception):
    """Carry a fixed public error without leaking paths or parser diagnostics."""

    def __init__(self, code: str, message: str) -> None:
        """Store a stable error code and safe, developer-authored message."""
        super().__init__(message)
        self.code = code
        self.message = message


def _guard_audit(event: str, args: tuple[Any, ...]) -> None:
    """Deny Python-mediated networking and child commands during local parsing."""
    if event.startswith("socket.") and event not in {
        "socket.__new__",
        "socket.gethostname",
    }:
        raise PermissionError("Parser network access is disabled")
    if event in {
        "subprocess.Popen",
        "os.system",
        "os.exec",
        "os.posix_spawn",
        "os.fork",
        "os.forkpty",
        "pty.spawn",
    }:
        raise PermissionError("Parser child processes are disabled")


def _apply_process_limits() -> None:
    """Install Unix resource limits before importing any document libraries."""
    try:
        import resource
        import signal

        # This child-owned deadline also kills a sleeping orphan if its parent dies.
        signal.signal(signal.SIGALRM, signal.SIG_DFL)
        signal.alarm(MAX_PARSER_SECONDS)
        limits = (
            (resource.RLIMIT_AS, MAX_MEMORY_BYTES),
            (resource.RLIMIT_CPU, MAX_PARSER_SECONDS),
            (resource.RLIMIT_FSIZE, MAX_RESULT_BYTES),
            (resource.RLIMIT_NOFILE, 128),
            (resource.RLIMIT_CORE, 0),
        )
        for limit_kind, maximum in limits:
            _, hard = resource.getrlimit(limit_kind)
            maximum = maximum if hard == resource.RLIM_INFINITY else min(maximum, hard)
            resource.setrlimit(limit_kind, (maximum, maximum))
    except (ImportError, OSError, ValueError) as error:
        raise ParseError(
            "limits_unavailable", "Parser resource limits are unavailable."
        ) from error
    for name in (
        "OMP_NUM_THREADS",
        "OPENBLAS_NUM_THREADS",
        "MKL_NUM_THREADS",
        "NUMEXPR_NUM_THREADS",
    ):
        os.environ[name] = "1"
    for name in ("HF_HUB_OFFLINE", "TRANSFORMERS_OFFLINE", "HF_HUB_DISABLE_TELEMETRY"):
        os.environ[name] = "1"
    sys.addaudithook(_guard_audit)


def _json_bytes(value: Any) -> bytes:
    """Encode deterministic compact UTF-8 JSON for accounting and output."""
    return json.dumps(
        value, ensure_ascii=False, allow_nan=False, separators=(",", ":")
    ).encode("utf-8")


def _new_result(extension: str) -> dict[str, Any]:
    """Create the endpoint's single result shape and private byte accounting."""
    return {
        "version": 1,
        "format": extension,
        "records": [],
        "index": {"pages": [], "sections": [], "sheets": []},
        "truncated": False,
        "truncation_reasons": [],
        "warnings": [],
        "_bytes": 0,
        "_index_keys": set(),
        "_section_keys": set(),
    }


def _truncate(result: dict[str, Any], reason: str) -> None:
    """Record a bounded, machine-readable explanation of omitted content."""
    result["truncated"] = True
    if reason not in result["truncation_reasons"]:
        result["truncation_reasons"].append(reason)


def _text(value: Any) -> str:
    """Normalize source values to safe Unicode without evaluating formulas."""
    return str(value).replace("\x00", "").encode("utf-8", "replace").decode("utf-8")


def _add_index(
    result: dict[str, Any], kind: str, key: str, value: str | int, text: str
) -> None:
    """Add at most one small preview per coordinate within the shared budget."""
    identity = (kind, value)
    if identity in result["_index_keys"]:
        return
    if len(result["_index_keys"]) >= MAX_INDEX_ENTRIES:
        _truncate(result, "index_limit")
        return
    entry = {key: value, "text": text[:MAX_INDEX_TEXT]}
    size = len(_json_bytes(entry)) + 1
    if result["_bytes"] + size > MAX_RESULT_BYTES - 65536:
        _truncate(result, "result_limit")
        return
    result["index"][kind].append(entry)
    result["_index_keys"].add(identity)
    result["_bytes"] += size


def _add_record(result: dict[str, Any], text: Any, **coordinates: Any) -> bool:
    """Append one bounded record, returning false when parsing must stop."""
    if len(result["records"]) >= MAX_RECORDS:
        _truncate(result, "record_limit")
        return False
    content = _text(text)
    if not content.strip():
        return True
    if len(content) > MAX_RECORD_TEXT:
        content = content[:MAX_RECORD_TEXT]
        _truncate(result, "record_text_limit")
    for key in ("sheet", "section"):
        if key in coordinates:
            label = _text(coordinates[key])
            if len(label) > MAX_INDEX_TEXT:
                _truncate(result, "label_limit")
            coordinates[key] = label[:MAX_INDEX_TEXT]
    section = coordinates.get("section")
    if section and section not in result["_section_keys"]:
        if len(result["_section_keys"]) >= MAX_SECTIONS:
            _truncate(result, "section_limit")
            return False
        result["_section_keys"].add(section)
    record = {"id": len(result["records"]) + 1, "text": content, **coordinates}
    size = len(_json_bytes(record)) + 1
    if result["_bytes"] + size > MAX_RESULT_BYTES - 65536:
        _truncate(result, "result_limit")
        return False
    result["records"].append(record)
    result["_bytes"] += size
    for kind, key in (("pages", "page"), ("sections", "section"), ("sheets", "sheet")):
        if key in coordinates:
            _add_index(result, kind, key, coordinates[key], content)
    return True


def _parse_text(path: Path, result: dict[str, Any]) -> None:
    """Read UTF-8 text incrementally, preserving multibyte character boundaries."""
    with path.open("r", encoding="utf-8-sig", errors="strict") as source:
        while content := source.read(MAX_RECORD_TEXT):
            if not _add_record(result, content, section="Text"):
                break


def _check_xlsx_archive(path: Path) -> None:
    """Reject oversized ZIP expansion, encryption, embedded macros, and XML DTDs."""
    import zipfile

    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if (
            len(entries) > MAX_ARCHIVE_MEMBERS
            or sum(item.file_size for item in entries) > MAX_ARCHIVE_BYTES
        ):
            raise ParseError("archive_limit", "Workbook exceeds safe archive limits.")
        names = {item.filename.lower() for item in entries}
        if any("vbaproject" in name or "macrosheets/" in name for name in names):
            raise ParseError(
                "unsupported_format", "Macro-enabled workbooks are not supported."
            )
        for entry in entries:
            if entry.flag_bits & 1 or entry.file_size > MAX_INPUT_BYTES:
                raise ParseError(
                    "archive_limit", "Workbook exceeds safe archive limits."
                )
            if entry.file_size > max(entry.compress_size, 1) * 1000:
                raise ParseError(
                    "archive_limit", "Workbook exceeds safe compression limits."
                )
            if not entry.filename.lower().endswith((".xml", ".rels")):
                continue
            with archive.open(entry) as source:
                previous = b""
                while chunk := source.read(65536):
                    scan = (previous + chunk).lower()
                    if b"<!doctype" in scan or b"<!entity" in scan:
                        raise ParseError(
                            "invalid_file",
                            "Workbook XML declarations are not supported.",
                        )
                    if (
                        entry.filename == "[Content_Types].xml"
                        and b"macroenabled" in scan
                    ):
                        raise ParseError(
                            "unsupported_format",
                            "Macro-enabled workbooks are not supported.",
                        )
                    previous = scan[-32:]


def _parse_xlsx(path: Path, result: dict[str, Any]) -> None:
    """Read cached XLSX values only, with bounded sheet, row, and cell traversal."""
    _check_xlsx_archive(path)
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(
        path, read_only=True, data_only=True, keep_links=False, keep_vba=False
    )
    result["warnings"].append(
        "Spreadsheet formulas are not evaluated; only saved cached values are read."
    )
    visited = 0
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            _truncate(result, "sheet_limit")
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            _add_index(result, "sheets", "sheet", sheet.title[:MAX_INDEX_TEXT], "")
            columns = min(sheet.max_column or MAX_COLUMNS, MAX_COLUMNS)
            rows = min(sheet.max_row or MAX_CELLS, MAX_CELLS)
            if (sheet.max_column or 0) > MAX_COLUMNS:
                _truncate(result, "column_limit")
            if (sheet.max_row or 0) > MAX_CELLS:
                _truncate(result, "row_limit")
            for row_number, row in enumerate(
                sheet.iter_rows(max_row=rows, max_col=columns), start=1
            ):
                for column_number, cell in enumerate(row, start=1):
                    if visited >= MAX_CELLS:
                        _truncate(result, "cell_limit")
                        return
                    visited += 1
                    if cell.value is not None and not _add_record(
                        result,
                        cell.value,
                        sheet=sheet.title,
                        section=sheet.title,
                        row=row_number,
                        column=column_number,
                        cell=f"{get_column_letter(column_number)}{row_number}",
                    ):
                        return
    finally:
        workbook.close()


def _check_image(path: Path) -> None:
    """Validate single-frame raster dimensions before any image decoding or OCR."""
    import warnings

    from PIL import Image

    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as image:
                if image.format not in {"PNG", "JPEG", "TIFF", "BMP", "WEBP"}:
                    raise ParseError(
                        "unsupported_format", "Unsupported raster image encoding."
                    )
                if (
                    image.width < 1
                    or image.height < 1
                    or image.width * image.height > MAX_IMAGE_PIXELS
                ):
                    raise ParseError("pixel_limit", "Image exceeds the pixel limit.")
                if getattr(image, "n_frames", 1) != 1:
                    raise ParseError(
                        "unsupported_format", "Multi-frame images are not supported."
                    )
                image.verify()
    except (Image.DecompressionBombWarning, Image.DecompressionBombError) as error:
        raise ParseError("pixel_limit", "Image exceeds the pixel limit.") from error


def _bounded_image_size(width: int, height: int) -> tuple[int, int]:
    """Fit even extreme aspect ratios within Docling's intermediate pixel cap."""
    scale = min(
        1.0,
        math.sqrt((MAX_IMAGE_PIXELS - 10000) / (OCR_RASTER_SCALE**2 * width * height)),
    )
    target_width, target_height = (
        max(1, int(width * scale)),
        max(1, int(height * scale)),
    )
    if target_width >= target_height:
        raster_width = MAX_IMAGE_PIXELS // math.ceil(target_height * OCR_RASTER_SCALE)
        target_width = min(
            target_width, max(1, math.floor(raster_width / OCR_RASTER_SCALE))
        )
    else:
        raster_height = MAX_IMAGE_PIXELS // math.ceil(target_width * OCR_RASTER_SCALE)
        target_height = min(
            target_height, max(1, math.floor(raster_height / OCR_RASTER_SCALE))
        )
    return target_width, target_height


def _check_pdf(path: Path) -> None:
    """Bound PDF page count and raster dimensions before Docling renders pages."""
    from pypdf import PdfReader

    reader = PdfReader(path, strict=True)
    if reader.is_encrypted:
        raise ParseError("encrypted_file", "Encrypted PDFs are not supported.")
    if len(reader.pages) > MAX_PAGES:
        raise ParseError("page_limit", "PDF exceeds the page limit.")
    for page in reader.pages:
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        unit = float(page.get("/UserUnit", 1))
        if not all(
            math.isfinite(value) and value > 0 for value in (width, height, unit)
        ):
            raise ParseError("invalid_file", "PDF has invalid page dimensions.")
        # Docling 2.60 renders OCR at 3x plus a 1.5x PDFium intermediate.
        if (
            math.ceil(width * unit * OCR_RASTER_SCALE)
            * math.ceil(height * unit * OCR_RASTER_SCALE)
            > MAX_IMAGE_PIXELS
        ):
            raise ParseError("pixel_limit", "PDF page exceeds the raster pixel limit.")
        crop_width, crop_height = float(page.cropbox.width), float(page.cropbox.height)
        if not all(
            math.isfinite(value) and value > 0 for value in (crop_width, crop_height)
        ):
            raise ParseError("invalid_file", "PDF has invalid crop dimensions.")
        if (
            math.ceil(crop_width * unit * OCR_RASTER_SCALE)
            * math.ceil(crop_height * unit * OCR_RASTER_SCALE)
            > MAX_IMAGE_PIXELS
        ):
            raise ParseError("pixel_limit", "PDF crop exceeds the raster pixel limit.")
        resources = page.get("/Resources")
        pending = [resources] if resources is not None else []
        checked: set[int] = set()
        while pending:
            resource = pending.pop().get_object()
            if id(resource) in checked:
                continue
            checked.add(id(resource))
            if len(checked) > 10000:
                raise ParseError("resource_limit", "PDF exceeds the resource limit.")
            for reference in (
                resource.get("/XObject", {}).get_object().values()
                if "/XObject" in resource
                else ()
            ):
                obj = reference.get_object()
                if obj.get("/Subtype") == "/Image":
                    image_width, image_height = (
                        int(obj.get("/Width", 0)),
                        int(obj.get("/Height", 0)),
                    )
                    if (
                        image_width < 1
                        or image_height < 1
                        or image_width * image_height > MAX_IMAGE_PIXELS
                    ):
                        raise ParseError(
                            "pixel_limit", "PDF image exceeds the pixel limit."
                        )
                elif obj.get("/Subtype") == "/Form" and "/Resources" in obj:
                    pending.append(obj["/Resources"])


def _parse_docling(path: Path, extension: str, result: dict[str, Any]) -> None:
    """Convert a local PDF or raster with explicitly offline, CPU-only Docling."""
    if extension == "pdf":
        _check_pdf(path)
    else:
        _check_image(path)
    artifacts = os.environ.get(DOCLING_ARTIFACTS_ENV, "")
    if (
        not artifacts
        or not Path(artifacts).is_absolute()
        or not Path(artifacts).is_dir()
    ):
        raise ParseError(
            "provider_unavailable",
            "PDF/image extraction requires provisioned local Docling models.",
        )
    rapid_models = (
        "onnx/PP-OCRv4/det/ch_PP-OCRv4_det_infer.onnx",
        "onnx/PP-OCRv4/cls/ch_ppocr_mobile_v2.0_cls_infer.onnx",
        "onnx/PP-OCRv4/rec/ch_PP-OCRv4_rec_infer.onnx",
        "paddle/PP-OCRv4/rec/ch_PP-OCRv4_rec_infer/ppocr_keys_v1.txt",
        "font.ttf",
    )
    if any(
        not (Path(artifacts) / "RapidOcr" / name).is_file() for name in rapid_models
    ):
        raise ParseError(
            "provider_unavailable", "Provisioned local OCR model files are missing."
        )
    try:
        from docling.datamodel.accelerator_options import (
            AcceleratorDevice,
            AcceleratorOptions,
        )
        from docling.datamodel.base_models import ConversionStatus, InputFormat
        from docling.datamodel.pipeline_options import (
            RapidOcrOptions,
            ThreadedPdfPipelineOptions,
        )
        from docling.document_converter import (
            DocumentConverter,
            ImageFormatOption,
            PdfFormatOption,
        )

        options = ThreadedPdfPipelineOptions(
            artifacts_path=Path(artifacts),
            enable_remote_services=False,
            allow_external_plugins=False,
            document_timeout=float(MAX_PARSER_SECONDS - 10),
            do_ocr=True,
            do_table_structure=True,
            generate_page_images=False,
            generate_picture_images=False,
            ocr_batch_size=1,
            layout_batch_size=1,
            table_batch_size=1,
            queue_max_size=1,
            accelerator_options=AcceleratorOptions(
                device=AcceleratorDevice.CPU, num_threads=1
            ),
            ocr_options=RapidOcrOptions(
                backend="onnxruntime",
                font_path=str(Path(artifacts) / "RapidOcr" / "font.ttf"),
                rec_font_path=str(Path(artifacts) / "RapidOcr" / "font.ttf"),
                rapidocr_params={
                    "EngineConfig.onnxruntime.intra_op_num_threads": 1,
                    "EngineConfig.onnxruntime.inter_op_num_threads": 1,
                    "EngineConfig.onnxruntime.use_cuda": False,
                    "EngineConfig.onnxruntime.use_dml": False,
                    "Rec.rec_keys_path": str(
                        Path(artifacts) / "RapidOcr" / rapid_models[3]
                    ),
                },
            ),
        )
        converter = DocumentConverter(
            allowed_formats=[InputFormat.PDF, InputFormat.IMAGE],
            format_options={
                InputFormat.PDF: PdfFormatOption(pipeline_options=options),
                InputFormat.IMAGE: ImageFormatOption(pipeline_options=options),
            },
        )
        converter.initialize_pipeline(
            InputFormat.PDF if extension == "pdf" else InputFormat.IMAGE
        )
    except (ImportError, OSError, RuntimeError, ValueError) as error:
        raise ParseError(
            "provider_unavailable",
            "Compatible Docling and local OCR models are unavailable.",
        ) from error
    conversion = converter.convert(
        path,
        max_num_pages=MAX_PAGES,
        max_file_size=MAX_INPUT_BYTES,
        raises_on_error=False,
    )
    if conversion.status not in {
        ConversionStatus.SUCCESS,
        ConversionStatus.PARTIAL_SUCCESS,
    }:
        raise ParseError("parse_failed", "Docling could not extract this document.")
    if conversion.status == ConversionStatus.PARTIAL_SUCCESS:
        _truncate(result, "docling_partial")
        result["warnings"].append(
            "Docling returned partial content; some pages or OCR regions may be missing."
        )
    document = conversion.document
    section = "Document"
    cells = 0
    for item, _level in document.iterate_items():
        label = str(getattr(item, "label", ""))
        text = getattr(item, "text", "")
        if label in {"section_header", "title"} or label.endswith(
            (".SECTION_HEADER", ".TITLE")
        ):
            section = _text(text)[:MAX_INDEX_TEXT] or "Document"
        provenance = getattr(item, "prov", [])
        coordinates: dict[str, Any] = {"section": section}
        if provenance:
            page = int(provenance[0].page_no)
            if page < 1 or page > MAX_PAGES:
                _truncate(result, "page_limit")
                continue
            coordinates["page"] = page
        table = getattr(item, "data", None)
        if table is not None and hasattr(table, "table_cells"):
            for cell in table.table_cells:
                if cells >= MAX_CELLS:
                    _truncate(result, "cell_limit")
                    return
                cells += 1
                if not _add_record(
                    result,
                    cell.text,
                    **coordinates,
                    row=cell.start_row_offset_idx + 1,
                    column=cell.start_col_offset_idx + 1,
                ):
                    return
        elif text and not _add_record(result, text, **coordinates):
            return


def parse_file(path: Path, extension: str) -> dict[str, Any]:
    """Parse one trusted staged local file into the bounded extraction contract.

    The production caller must use the CLI subprocess, which applies limits
    before importing parsing libraries. Direct calls are only useful for tests.
    """
    if extension not in SUPPORTED_EXTENSIONS:
        raise ParseError(
            "unsupported_format",
            "Supported formats are TXT, XLSX, PDF, and single-frame raster images.",
        )
    if not path.is_absolute() or path.is_symlink():
        raise ParseError("invalid_file", "A staged regular file is required.")
    info = path.stat()
    if not stat.S_ISREG(info.st_mode) or not 0 < info.st_size <= MAX_INPUT_BYTES:
        raise ParseError("file_limit", "File is empty or exceeds the input size limit.")
    result = _new_result(extension)
    if extension == "txt":
        _parse_text(path, result)
    elif extension == "xlsx":
        _parse_xlsx(path, result)
    elif extension == "pdf":
        _parse_docling(path, extension, result)
    else:
        import tempfile

        from PIL import Image

        _check_image(path)
        with Image.open(path) as source:
            if (
                math.ceil(source.width * OCR_RASTER_SCALE)
                * math.ceil(source.height * OCR_RASTER_SCALE)
                <= MAX_IMAGE_PIXELS
            ):
                _parse_docling(path, extension, result)
            else:
                # Include Docling's 1.5x intermediate over RapidOCR's 3x raster.
                target = _bounded_image_size(source.width, source.height)
                with tempfile.TemporaryDirectory(
                    prefix="extraction-image-"
                ) as directory:
                    reduced = source.resize(target, Image.Resampling.LANCZOS)
                    # RGB keeps the bounded temporary PNG below RLIMIT_FSIZE.
                    rgb = Image.new("RGB", reduced.size, "white")
                    if "A" in reduced.getbands() or "transparency" in reduced.info:
                        rgba = reduced.convert("RGBA")
                        rgb.paste(rgba, mask=rgba.getchannel("A"))
                        rgba.close()
                    else:
                        rgb.paste(reduced)
                    reduced.close()
                    reduced = rgb
                    raster = Path(directory) / "image.png"
                    try:
                        reduced.save(raster, format="PNG")
                    finally:
                        reduced.close()
                    _truncate(result, "image_downsampled")
                    result["warnings"].append(
                        "Image resolution was reduced to keep OCR within the raster pixel limit."
                    )
                    _parse_docling(raster, "png", result)
    for key in ("_bytes", "_index_keys", "_section_keys"):
        result.pop(key)
    if len(_json_bytes(result)) > MAX_RESULT_BYTES:
        raise ParseError("result_limit", "Extraction exceeds the result size limit.")
    return result


def main(arguments: list[str] | None = None) -> int:
    """Apply child limits and print one sanitized JSON result or fixed error."""
    arguments = sys.argv[1:] if arguments is None else arguments
    try:
        _apply_process_limits()
        if len(arguments) != 2:
            raise ParseError(
                "invalid_input", "Expected a staged file path and extension."
            )
        # Native libraries may write to fd 1; preserve it only for final JSON.
        output = os.dup(sys.stdout.fileno())
        with open(os.devnull, "wb") as discard:
            os.dup2(discard.fileno(), sys.stdout.fileno())
            os.dup2(discard.fileno(), sys.stderr.fileno())
        result = parse_file(Path(arguments[0]), arguments[1])
        exit_code = 0
    except ParseError as error:
        result = {"error": {"code": error.code, "message": error.message}}
        exit_code = 1
    except UnicodeError:
        result = {
            "error": {
                "code": "invalid_encoding",
                "message": "TXT files must contain valid UTF-8 text.",
            }
        }
        exit_code = 1
    except MemoryError:
        result = {
            "error": {
                "code": "resource_limit",
                "message": "Parser exceeded its memory limit.",
            }
        }
        exit_code = 1
    except Exception:  # noqa: BLE001 - subprocess boundary must not expose parser details.
        result = {
            "error": {
                "code": "parse_failed",
                "message": "File could not be safely parsed.",
            }
        }
        exit_code = 1
    payload = _json_bytes(result)
    with (
        os.fdopen(output, "wb")
        if "output" in locals()
        else sys.stdout.buffer as destination
    ):
        destination.write(payload)
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
